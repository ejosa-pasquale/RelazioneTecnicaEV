\
import datetime as dt
import io
from pathlib import Path

import streamlit as st
from streamlit_drawable_canvas import st_canvas

from docx import Document

def iter_all_paragraphs_doc(doc: Document):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
                for nt in cell.tables:
                    for nrow in nt.rows:
                        for ncell in nrow.cells:
                            for np in ncell.paragraphs:
                                yield np
    for section in doc.sections:
        for hf in [section.header, section.footer]:
            for p in hf.paragraphs:
                yield p
            for t in hf.tables:
                for row in t.rows:
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            yield p

_AGG_RE = re.compile(r"\baggiungere\s+([^\n\r]+)", re.IGNORECASE)

def extract_aggiungere_fields(template_bytes: bytes):
    """
    Trova occorrenze del tipo 'aggiungere XXX' nel template (anche in tabelle/header/footer).
    Ritorna una lista di dict: {key,label,token}.
    - token: stringa completa da sostituire (es. 'aggiungere DESCRIZIONE INTERVENTO')
    - key: chiave normalizzata per Streamlit
    """
    d = Document(io.BytesIO(template_bytes))
    found = {}
    for p in iter_all_paragraphs_doc(d):
        txt = p.text or ""
        for m in _AGG_RE.finditer(txt):
            label = m.group(1).strip()
            token = txt[m.start():m.end()]  # substring matched
            # Normalizza key
            key = re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")
            if not key:
                key = f"campo_{len(found)+1}"
            # se lo stesso label appare più volte, unifica
            found.setdefault(key, {"key": key, "label": label, "token": f"aggiungere {label}"})
    return list(found.values())


from generator import (
    AllegatoItem,
    ColonninaItem,
    EsecutriceData,
    PhotoItem,
    ProgettistaData,
    RelazioneData,
    generate_document,
    prepare_template,
)

APP_DIR = Path(__file__).parent
DEFAULT_TEMPLATE_PATH = APP_DIR / "templates" / "template_aggiungere.dotx"

st.set_page_config(page_title="Relazione tecnica EV", layout="wide")
st.title("Relazione tecnica — Generatore veloce")
st.caption("Flusso consigliato: 1) carica template 2) **Prepara template (agganci nel corpo)** 3) usa sempre quel template preparato 4) genera DOCX.")

if "template_bytes" not in st.session_state:
    st.session_state.template_bytes = DEFAULT_TEMPLATE_PATH.read_bytes()
    st.session_state.aggiungere_fields = extract_aggiungere_fields(st.session_state.template_bytes)
if "photos" not in st.session_state:
    st.session_state.photos = []
if "diagram" not in st.session_state:
    st.session_state.diagram = None
if "colonnine" not in st.session_state:
    st.session_state.colonnine = []
if "allegati" not in st.session_state:
    st.session_state.allegati = []
if "aggiungere_fields" not in st.session_state:
    st.session_state.aggiungere_fields = []
if "aggiungere_values" not in st.session_state:
    st.session_state.aggiungere_values = {}

with st.sidebar:
    st.header("1) Template")
    up = st.file_uploader("Carica template DOCX (opzionale)", type=["docx"])
    if up:
        st.session_state.template_bytes = up.getvalue()
        st.success("Template caricato.")
        st.session_state.aggiungere_fields = extract_aggiungere_fields(st.session_state.template_bytes)

    colA, colB = st.columns(2)
    with colA:
        if st.button("Prepara template (agganci)", use_container_width=True):
            st.session_state.template_bytes = prepare_template(st.session_state.template_bytes)
            st.success("Template preparato: punti di aggancio inseriti nel corpo testo.")
    with colB:
        st.download_button(
            "Scarica template preparato",
            data=st.session_state.template_bytes,
            file_name="template_preparato.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    st.divider()
    st.header("2) Progettista (fisso)")
    progettista = ProgettistaData(
        nome=st.text_input("Nome", "Ing. Pasquale Senese"),
        indirizzo=st.text_input("Indirizzo", "Via Francesco Soave 30 - 20135 Milano (MI)"),
        cell=st.text_input("Cell", "340 5731381"),
        email=st.text_input("Email", "pasquale.senese@ingpec.eu"),
        piva=st.text_input("P.IVA", "14572980960"),
    )

    st.divider()
    st.header("3) Ditta esecutrice (opzionale)")
    esec_nome = st.text_input("Nome/Ragione sociale", "")
    esec_ind = st.text_input("Indirizzo", "")
    esec_piva = st.text_input("P.IVA", "")
    esecutrice = EsecutriceData(esec_nome, esec_ind, esec_piva) if (esec_nome or esec_ind or esec_piva) else None

tab1, tab2, tab3, tab4 = st.tabs(["Dati progetto", "Colonnine", "Foto & Diagramma", "Schede tecniche"])

with tab1:
    st.subheader("Dati progetto")
    c1, c2, c3 = st.columns(3)
    with c1:
        luogo = st.text_input("Luogo", "Busnago")
        data = st.date_input("Data", dt.date.today())
        luogo_data = f"{luogo}, {data.strftime('%d/%m/%Y')}"
        committente = st.text_input("Committente", "Nome")
    with c2:
        sito_ind = st.text_input("Indirizzo sito", "via della SS. Annunziata 32/A")
        sito_cc = st.text_input("CAP e Città", "55100 Lucca")
        oggetto = st.text_input("Oggetto", "nuovo impianto di ricarica per veicolo elettrico")
    with c3:
        distanza = st.number_input("Distanza (m)", min_value=0, value=60, step=1)
        potenza = st.number_input("Potenza impegnata (kW)", min_value=0.0, value=4.0, step=0.5)

    
    st.subheader("Campi del template (AGGIUNGERE …)")
    st.caption("Nel template hai scritto 'aggiungere XXX': qui inserisci il valore che sostituirà quel testo nel corpo del documento.")
    if not st.session_state.aggiungere_fields:
        st.info("Nessun campo 'aggiungere …' trovato nel template.")
    else:
        for f in st.session_state.aggiungere_fields:
            k = f"aggi_{f['key']}"
            default = st.session_state.aggiungere_values.get(f['key'], "")
            val = st.text_area(f['label'], value=default, height=80, key=k)
            st.session_state.aggiungere_values[f['key']] = val

    st.subheader("Layout d'impianto (testo nel corpo)")
    layout_incluso = st.text_area("Incluso (1 riga = 1 bullet)", height=120)
    layout_escluso = st.text_area("Escluso (1 riga = 1 bullet)", height=120)

    st.subheader("Dati elettrici")
    c4, c5, c6 = st.columns(3)
    with c4:
        ik_t = st.number_input("Ik trifase (kA)", min_value=0.0, value=10.0, step=0.5)
    with c5:
        ik_m = st.number_input("Ik monofase (kA)", min_value=0.0, value=6.0, step=0.5)
    with c6:
        cavo_len = st.number_input("Lunghezza cavo (m)", min_value=0, value=60, step=1)
        cavo_tipo = st.text_input("Tipo cavo", "FG16OM16 3G6 0.6/1kV")

    data_obj = RelazioneData(
        luogo_data=luogo_data,
        committente=committente,
        sito_indirizzo=sito_ind,
        sito_cap_citta=sito_cc,
        oggetto=oggetto,
        distanza_m=int(distanza),
        potenza_impegnata_kw=float(potenza),
        ik_trifase_ka=float(ik_t),
        ik_monofase_ka=float(ik_m),
        cavo_lunghezza_m=int(cavo_len),
        cavo_tipo=cavo_tipo,
        layout_incluso=layout_incluso,
        layout_escluso=layout_escluso,
    )

with tab2:
    st.subheader("Colonnine")
    cc1, cc2 = st.columns([3,1])
    with cc1:
        descr = st.text_input("Descrizione (modello/potenza/connettori/note)", "")
    with cc2:
        qty = st.number_input("Qtà", min_value=1, value=1, step=1)

    if st.button("Aggiungi colonnina", use_container_width=True):
        if descr.strip():
            st.session_state.colonnine.append(ColonninaItem(descrizione=descr.strip(), quantita=int(qty)))
            st.success("Aggiunta.")
        else:
            st.warning("Inserisci una descrizione.")

    if st.session_state.colonnine:
        st.write("Elenco colonnine:")
        for i, item in enumerate(list(st.session_state.colonnine)):
            r1, r2, r3 = st.columns([0.8, 3, 0.8])
            with r1:
                st.write(f"n. {item.quantita}")
            with r2:
                st.session_state.colonnine[i].descrizione = st.text_input("Descrizione", item.descrizione, key=f"col_{i}")
            with r3:
                if st.button("Rimuovi", key=f"col_rm_{i}"):
                    st.session_state.colonnine.pop(i)
                    st.rerun()

with tab3:
    st.subheader("Foto")
    ups = st.file_uploader("Carica foto (JPG/PNG)", type=["jpg","jpeg","png"], accept_multiple_files=True)
    if ups:
        existing = {p.filename for p in st.session_state.photos}
        for u in ups:
            if u.name not in existing:
                st.session_state.photos.append(PhotoItem(filename=u.name, content=u.getvalue(), caption=""))
        st.success("Foto aggiunte.")

    if st.session_state.photos:
        for i, ph in enumerate(list(st.session_state.photos)):
            a, b, c = st.columns([1,2,0.6])
            with a:
                st.image(ph.content, use_container_width=True)
            with b:
                st.session_state.photos[i].caption = st.text_input("Didascalia", ph.caption, key=f"cap_{i}")
            with c:
                if st.button("Rimuovi", key=f"ph_rm_{i}"):
                    st.session_state.photos.pop(i); st.rerun()

    st.divider()
    st.subheader("Diagramma impianto")
    mode = st.radio("Metodo", ["Carica immagine", "Disegna"], horizontal=True)
    if mode == "Carica immagine":
        d = st.file_uploader("Carica diagramma (JPG/PNG)", type=["jpg","jpeg","png"])
        if d:
            st.session_state.diagram = d.getvalue()
            st.image(st.session_state.diagram, use_container_width=True)
    else:
        canvas = st_canvas(
            fill_color="rgba(255,255,255,0)",
            stroke_width=3,
            stroke_color="#000000",
            background_color="#FFFFFF",
            height=420,
            drawing_mode="freedraw",
            key="canvas",
        )
        if st.button("Salva disegno come diagramma", type="primary"):
            if canvas.image_data is not None:
                import PIL.Image
                img = PIL.Image.fromarray(canvas.image_data.astype("uint8"), mode="RGBA").convert("RGB")
                buf = io.BytesIO()
                img.save(buf, format="PNG")
                st.session_state.diagram = buf.getvalue()
                st.success("Diagramma salvato.")
        if st.session_state.diagram:
            st.image(st.session_state.diagram, caption="Diagramma pronto", use_container_width=True)

with tab4:
    st.subheader("Schede tecniche")
    att = st.file_uploader("Carica schede tecniche (PDF/DOCX)", type=["pdf","docx"], accept_multiple_files=True)
    if att:
        existing = {a.filename for a in st.session_state.allegati}
        for u in att:
            if u.name in existing:
                continue
            kind = "pdf" if u.name.lower().endswith(".pdf") else "docx"
            st.session_state.allegati.append(AllegatoItem(filename=u.name, content=u.getvalue(), kind=kind))
        st.success("Allegati aggiunti.")
    if st.session_state.allegati:
        for i, a in enumerate(list(st.session_state.allegati)):
            x, y = st.columns([4,1])
            with x:
                st.write(f"📎 {a.filename} ({a.kind})")
            with y:
                if st.button("Rimuovi", key=f"att_rm_{i}"):
                    st.session_state.allegati.pop(i); st.rerun()

st.divider()
st.subheader("Generazione")
with st.expander("Debug template"):
    from docx import Document
    import io as _io
    try:
        _d = Document(_io.BytesIO(st.session_state.template_bytes))
        text = "\n".join([p.text for p in _d.paragraphs])
        st.write("Marker presenti nel body:", {m: (m in text) for m in ["{{LAYOUT_DESCRITTIVO}}","{{COLONNINE}}","{{FOTO_GALLERY}}","{{DIAGRAMMA_IMPIANTO}}","{{ALLEGATI_SCHEDA_TECNICA}}","{{DITTA_ESECUTRICE}}"]})
        st.caption("Nota: se i titoli/marker sono dentro tabelle, non compaiono qui. La v12 li cerca anche dentro le tabelle durante la generazione.")
    except Exception as e:
        st.write(e)

fname = st.text_input("Nome file", "relazione_tecnica")
colG1, colG2 = st.columns([1,1])
with colG1:
    if st.button("Genera DOCX", type="primary", use_container_width=True):
        docx = generate_document(
            template=st.session_state.template_bytes,
            data=data_obj,
            progettista=progettista,
            esecutrice=esecutrice,
            colonnine=st.session_state.colonnine,
            photos=st.session_state.photos,
            diagram_bytes=st.session_state.diagram,
            allegati=st.session_state.allegati,
            extra_fields={f['label']: st.session_state.aggiungere_values.get(f['key'], '') for f in st.session_state.aggiungere_fields},
        )
        st.session_state["last_docx"] = docx
        st.success("Documento generato.")

with colG2:
    if st.session_state.get("last_docx"):
        st.download_button(
            "Scarica DOCX",
            data=st.session_state["last_docx"],
            file_name=f"{fname}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )